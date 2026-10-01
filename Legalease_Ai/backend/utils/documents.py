from pathlib import Path
from html import escape

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet


def safe_filename(name: str) -> str:
    """Create a safe filename."""
    invalid_chars = '<>:"/\\|?*'
    for char in invalid_chars:
        name = name.replace(char, "_")

    name = name.strip()

    if not name:
        name = "generated_document"

    return name


def format_txt(text: str) -> bytes:
    """Convert text to TXT bytes."""
    return text.encode("utf-8")


def format_docx(text: str) -> bytes:
    """Convert text to DOCX bytes."""
    document = Document()

    for line in text.splitlines():
        if line.strip():
            document.add_paragraph(line)
        else:
            document.add_paragraph("")

    output_path = Path("generated_document.docx")
    document.save(output_path)

    data = output_path.read_bytes()
    output_path.unlink(missing_ok=True)

    return data


def format_pdf(text: str) -> bytes:
    """Convert text to PDF bytes."""
    output_path = Path("generated_document.pdf")

    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        rightMargin=50,
        leftMargin=50,
        topMargin=50,
        bottomMargin=50,
    )

    styles = getSampleStyleSheet()
    story = []

    for line in text.splitlines():
        if line.strip():
            story.append(Paragraph(escape(line), styles["BodyText"]))
            story.append(Spacer(1, 8))
        else:
            story.append(Spacer(1, 8))

    doc.build(story)

    data = output_path.read_bytes()
    output_path.unlink(missing_ok=True)

    return data


def text_to_html(text: str) -> str:
    """Convert plain text into simple HTML preview."""
    escaped = escape(text)
    return escaped.replace("\n", "<br>")